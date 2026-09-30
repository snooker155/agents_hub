"""Slides views: the spec's layouts and themes, the agent tools that build a
deck, the .pptx export and its download route.

The renderer (dashboard/frontend/src/views/slides/) and the export
(views/slides_pptx.py) draw the same layouts in the same palette; the palette
file is shared by copy, and the test below keeps the copy honest.
"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import pytest

from views.models import SLIDE_LAYOUTS, SLIDE_THEMES, ViewValidationError, validate_spec

ROOT = Path(__file__).resolve().parents[1]


_tokens: list = []


def _bind_view(vid):
    from common.agent_context import current_view_id
    _tokens.append(current_view_id.set(vid))


@pytest.fixture(autouse=True)
def _unbind_view():
    """The active view is a contextvar; left set, it leaks into later tests."""
    yield
    from common.agent_context import current_view_id
    while _tokens:
        current_view_id.reset(_tokens.pop())


def _call(tool, **kwargs) -> dict:
    return json.loads(tool.invoke(kwargs))


FULL_DECK = {
    "theme": "ocean",
    "accent": "#0f766e",
    "footer": "Team · 2026",
    "slides": {
        "a": {"order": 0, "layout": "title", "title": "What is new", "subtitle": "September", "icon": "🚀"},
        "b": {"order": 1, "layout": "section", "title": "Part one"},
        "c": {"order": 2, "layout": "content", "title": "Deploy", "body": "- **compose**\n  - nested\n1. one\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n```\ncode\n```", "notes": "Say it slowly."},
        "d": {"order": 3, "layout": "two_column", "title": "Before and after", "columns": ["### Before\n- manual", "### After\n- one click"]},
        "e": {"order": 4, "layout": "stats", "title": "Numbers", "items": [{"value": "8", "title": "features", "text": "deploy"}]},
        "f": {"order": 5, "layout": "cards", "title": "Cards", "items": [{"icon": "📦", "title": f"c{i}", "text": "x"} for i in range(6)]},
        "g": {"order": 6, "layout": "timeline", "title": "Stages", "items": [{"value": f"{i} Sep", "title": f"s{i}"} for i in range(8)]},
        "h": {"order": 7, "layout": "quote", "body": "All in one place.", "subtitle": "the team"},
        "i": {"order": 8, "layout": "image_left", "title": "Left", "image": "asset://shot.png", "caption": "cap", "body": "text"},
        "j": {"order": 9, "layout": "image_right", "title": "Right", "image": "shot.webp", "body": "text"},
        "k": {"order": 10, "layout": "image_full", "title": "Thanks", "image": "https://example.com/x.png"},
    },
}


# ── spec ─────────────────────────────────────────────────────────────────────

def test_the_frontend_palette_is_the_backend_palette():
    front = json.loads((ROOT / "dashboard/frontend/src/views/slideThemes.json").read_text(encoding="utf-8"))
    assert front == SLIDE_THEMES


def test_every_layout_validates():
    spec = validate_spec("slides", FULL_DECK)
    assert {s["layout"] for s in spec["slides"].values()} == set(SLIDE_LAYOUTS)
    assert spec["theme"] == "ocean" and spec["footer"] == "Team · 2026"


@pytest.mark.parametrize("slide,words", [
    ({"layout": "stats", "title": "x", "items": []}, "1 to 4 items"),
    ({"layout": "cards", "title": "x", "items": [{"title": "a"}] * 7}, "1 to 6 items"),
    ({"layout": "timeline", "title": "x", "items": [{"text": "no title or value"}]}, "title or a value"),
    ({"layout": "two_column", "title": "x", "columns": ["only one"]}, "exactly two"),
    ({"layout": "image_left", "title": "x"}, "needs an image"),
    ({"layout": "quote", "title": "x"}, "quotation in body"),
    ({"layout": "section"}, "needs a title"),
    ({"layout": "hero", "title": "x"}, "unknown layout"),
    ({"title": "x", "icon": "rocket"}, "single emoji"),
    ({"title": "x", "accent": "blue"}, "#rrggbb"),
    ({"title": "x", "items": [{"title": "a", "extra": 1}]}, "Extra inputs"),
    ({"title": "x", "bullets": ["a"]}, "not bullets"),
])
def test_a_slide_its_layout_cannot_draw_is_refused_with_the_reason(slide, words):
    with pytest.raises(ViewValidationError) as exc:
        validate_spec("slides", {"slides": {"s1": slide}})
    assert words in str(exc.value)
    assert "slide s1" in str(exc.value), "the error must name the slide"


def test_an_unknown_theme_is_refused():
    with pytest.raises(ViewValidationError, match="unknown theme"):
        validate_spec("slides", {"slides": {}, "theme": "neon"})


# ── tools ────────────────────────────────────────────────────────────────────

def test_slides_add_builds_layouts_and_replaces_by_id():
    from tools.views import slides_add
    from views.store import create_live_view, get_view
    vid = create_live_view("slides", "Deck").view_id
    _bind_view(vid)
    first = _call(slides_add, title="Numbers", layout="stats",
                  items=json.dumps([{"value": "8", "title": "features"}]), slide_id="n")
    assert first["ok"] and first["slide_id"] == "n" and first["replaced"] is False
    assert _call(slides_add, title="Compare", layout="two_column", columns=["a", "b"])["ok"]
    again = _call(slides_add, title="Numbers v2", layout="stats",
                  items=[{"value": "9", "title": "features"}], slide_id="n")
    assert again["replaced"] is True
    slides = get_view(vid)["spec"]["slides"]
    assert slides["n"]["title"] == "Numbers v2" and slides["n"]["order"] == 0
    assert sorted(s["order"] for s in slides.values()) == [0, 1]


def test_slides_add_refuses_a_bad_slide_and_changes_nothing():
    from tools.views import slides_add
    from views.store import create_live_view, get_view
    vid = create_live_view("slides", "Deck").view_id
    _bind_view(vid)
    res = _call(slides_add, title="Numbers", layout="stats")
    assert res["ok"] is False and "1 to 4 items" in res["error"]
    assert get_view(vid)["spec"]["slides"] == {}


def test_slides_add_keeps_a_deck_created_as_a_list():
    """A path op into a list would replace it with a one-slide map."""
    from tools.views import create_view_tools, slides_add
    from views.store import get_view
    create_view = create_view_tools(None)[0]
    made = _call(create_view, view_kind="slides", title="Deck", summary="s",
                 spec=json.dumps({"slides": [{"title": "One", "body": "x"}, {"title": "Two", "body": "y"}]}))
    _bind_view(made["view_id"])
    assert _call(slides_add, title="Three", body="z")["ok"]
    slides = get_view(made["view_id"])["spec"]["slides"]
    assert [s["title"] for s in sorted(slides.values(), key=lambda s: s["order"])] == ["One", "Two", "Three"]


def test_slides_style_sets_the_deck_look():
    from tools.views import slides_style
    from views.store import create_live_view, get_view
    vid = create_live_view("slides", "Deck").view_id
    _bind_view(vid)
    assert _call(slides_style, theme="Dark", accent="#ff0066", footer="Q3", numbers="false")["ok"]
    spec = get_view(vid)["spec"]
    assert (spec["theme"], spec["accent"], spec["footer"], spec["numbers"]) == ("dark", "#ff0066", "Q3", False)
    assert _call(slides_style, theme="neon")["ok"] is False
    assert get_view(vid)["spec"]["theme"] == "dark"
    assert _call(slides_style)["ok"] is False


def test_slides_export_writes_a_pptx_into_the_workspace(tmp_path):
    from pptx import Presentation

    from tools.views import create_view_tools, slides_add
    tools = {t.name: t for t in create_view_tools(str(tmp_path))}
    made = _call(tools["create_view"], view_kind="slides", title="Что нового", summary="s", spec='{"slides": {}}')
    _bind_view(made["view_id"])
    _call(slides_add, title="Cover", layout="title", subtitle="sub")
    _call(slides_add, title="Body", body="- a\n- b", notes="n1")
    out = _call(tools["slides_export"])
    assert out["ok"], out
    assert out["path"] == "presentations/Что-нового.pptx"
    prs = Presentation(str(tmp_path / out["path"]))
    assert len(prs.slides) == 2
    assert prs.slides[1].notes_slide.notes_text_frame.text == "n1"
    bad = _call(tools["slides_export"], path="../escape.pptx")
    assert bad["ok"] is False and "escapes" in bad["error"]


def test_slides_export_refuses_another_kind(tmp_path):
    from tools.views import create_view_tools
    tools = {t.name: t for t in create_view_tools(str(tmp_path))}
    made = _call(tools["create_view"], view_kind="markdown", title="M", summary="s", spec='{"markdown": "# hi"}')
    res = _call(tools["slides_export"], view_id=made["view_id"])
    assert res["ok"] is False and "not slides" in res["error"]


# ── the .pptx ────────────────────────────────────────────────────────────────

def _texts(slide) -> str:
    out = []
    for shp in slide.shapes:
        if shp.has_text_frame:
            out.append(shp.text_frame.text)
        if getattr(shp, "has_table", False) and shp.has_table:
            out.extend(c.text for r in shp.table.rows for c in r.cells)
    return "\n".join(out)


def test_every_layout_exports_with_its_content(tmp_path):
    from PIL import Image
    from pptx import Presentation

    from views.slides_pptx import build_pptx
    png = tmp_path / "shot.png"
    Image.new("RGB", (400, 300), "#4f7cac").save(png)
    webp = tmp_path / "shot.webp"
    Image.new("RGB", (300, 400), "#f6c453").save(webp, format="WEBP")
    view = {"view_id": "", "title": "Deck", "spec": validate_spec("slides", FULL_DECK)}
    data, warnings = build_pptx(view, with_warnings=True,
                                resolve_asset=lambda n: {"shot.png": png, "shot.webp": webp}.get(n))
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 11
    assert prs.slide_width == 12192000 and prs.slide_height == 6858000  # 16:9, 13.333 x 7.5 in
    texts = [_texts(s) for s in prs.slides]
    assert "What is new" in texts[0] and "September" in texts[0] and "🚀" in texts[0]
    assert "01" in texts[1] and "Part one" in texts[1]
    assert "compose" in texts[2] and "Say it slowly." == prs.slides[2].notes_slide.notes_text_frame.text
    assert "1" in texts[2] and "b" in texts[2]  # the table
    assert "Before" in texts[3] and "one click" in texts[3]
    assert "8" in texts[4] and "features" in texts[4]
    assert all(f"c{i}" in texts[5] for i in range(6))
    assert all(f"{i} Sep" in texts[6] for i in range(8))
    assert "All in one place." in texts[7] and "— the team" in texts[7]
    pictures = [sum(1 for sh in s.shapes if sh.shape_type == 13) for s in prs.slides]
    assert pictures[8] == 1 and pictures[9] == 1, "the png and the (converted) webp are embedded"
    assert pictures[10] == 0 and warnings == ["slide 11: remote image not embedded (https://example.com/x.png)"]
    assert "Team · 2026" in texts[2] and "3 / 11" in texts[2]


def test_long_text_is_shrunk_to_fit():
    from views.slides_pptx import blocks_height, fit_size, parse_markdown
    long = "\n".join(f"- point number {i} with quite a few words in it to wrap" for i in range(18))
    blocks = parse_markdown(long)
    size = fit_size(blocks, 1120, 480, 26, 13)
    assert size < 26 and blocks_height(blocks, size, 1120) <= 480 or size == 13


def test_markdown_parse_matches_the_renderer_rules():
    from views.slides_pptx import inline_runs, parse_markdown
    kinds = [b[0] for b in parse_markdown(
        "# T\n\n- one\n  - nested\n1. first\n2. second\n\n> q\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n```\nc\n```\n---\ntext\nmore")]
    assert kinds == ["heading", "bullet", "bullet", "bullet", "bullet", "quote", "table", "code", "hr", "para"]
    assert inline_runs("a **b** `c` [d](https://x.y)") == [
        ("a ", {}), ("b", {"bold": True}), (" ", {}), ("c", {"code": True}), (" ", {}), ("d", {"link": "https://x.y"})]


def test_pptx_filename_keeps_unicode_and_drops_separators():
    from views.slides_pptx import pptx_filename
    assert pptx_filename("Что нового: 2026/09") == "Что-нового-2026-09.pptx"
    assert pptx_filename("") == "slides.pptx"


# ── route ────────────────────────────────────────────────────────────────────

@pytest.fixture
def client():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    sys.path.insert(0, str(ROOT / "dashboard" / "backend"))
    from routes import views as views_routes
    app = FastAPI()
    app.include_router(views_routes.router)
    return TestClient(app)


def test_export_route_downloads_a_pptx(client):
    from pptx import Presentation

    from views.store import create_view
    env = create_view("slides", "Что нового", {"slides": {"a": {"title": "One", "body": "- x"}}}, workspace="wsX", summary="s")
    r = client.get(f"/api/views/{env.view_id}/export/pptx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml")
    assert "filename*=UTF-8''%D0%A7%D1%82%D0%BE-%D0%BD%D0%BE%D0%B2%D0%BE%D0%B3%D0%BE.pptx" in r.headers["content-disposition"]
    assert r.headers["x-export-warnings"] == "0"
    assert len(Presentation(io.BytesIO(r.content)).slides) == 1


def test_export_route_refuses_other_kinds_and_unknown_views(client):
    from views.store import create_view
    env = create_view("markdown", "M", {"markdown": "# hi"}, workspace="wsX", summary="s")
    assert client.get(f"/api/views/{env.view_id}/export/pptx").status_code == 400
    assert client.get("/api/views/vw_nope/export/pptx").status_code == 404
