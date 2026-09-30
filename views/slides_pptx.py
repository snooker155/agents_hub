"""
A slides view as a PowerPoint file.

``build_pptx(view)`` turns a ``slides`` view (views/models.py ``SlidesSpec``)
into .pptx bytes with python-pptx. It draws the same layouts on the same 16:9
grid and with the same palette (views/slide_themes.json) as the browser
renderer (dashboard/frontend/src/views/renderers/SlidesView.jsx), so what the
user saw is what they download: a cover on the hero gradient, section
dividers, markdown bodies with real bullets, two columns, images around or
under the text, quotes, big numbers, cards and a timeline, plus the footer,
slide numbers and speaker notes.

Geometry is written in the renderer's pixels (1280 x 720) and converted here;
one pixel is 9525 EMU at 96 dpi, which is exactly 13.333 x 7.5 inches.

Text is sized to fit before it is placed. PowerPoint only recomputes
"shrink text on overflow" when the text is edited, so a box that relied on it
would open overflowing; the size is estimated from the text instead (a
character is about half an em wide) and the box also carries normAutofit for
the editors that do honour it.

Images come only from the view's own assets. A remote URL is not fetched on
the server (it would be a request to an address the agent chose); the slide
gets a placeholder with the link and the export reports it in ``warnings``.
"""
from __future__ import annotations

import io
import math
import re
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple

from views.models import SLIDE_THEMES

PX = 9525  # EMU per pixel at 96 dpi
FONT = "Arial"  # on every machine that opens a .pptx, Latin and Cyrillic alike
W, H = 1280, 720

_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|!?\[[^\]]*\]\([^)]+\))")
_BULLET = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")


# ── markdown ─────────────────────────────────────────────────────────────────

def parse_markdown(md: str) -> List[Tuple]:
    """Blocks of the markdown subset slides use.

    ``("heading", level, text)``, ``("bullet", level, ordered, number, text)``,
    ``("para", text)``, ``("quote", text)``, ``("code", text)``,
    ``("table", rows)`` and ``("hr",)``. Anything else is a paragraph, so no
    text is ever dropped.
    """
    lines = (md or "").replace("\r\n", "\n").split("\n")
    blocks: List[Tuple] = []
    para: List[str] = []
    counters: Dict[int, int] = {}

    def flush() -> None:
        if para:
            blocks.append(("para", " ".join(s.strip() for s in para)))
            para.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            body = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                body.append(lines[i])
                i += 1
            blocks.append(("code", "\n".join(body)))
            i += 1
            continue
        if not stripped:
            flush()
            counters.clear()
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            flush()
            blocks.append(("heading", len(m.group(1)), m.group(2).strip()))
            i += 1
            continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", stripped):
            flush()
            blocks.append(("hr",))
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]):
            flush()
            rows = [_cells(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_cells(lines[i].strip()))
                i += 1
            blocks.append(("table", rows))
            continue
        m = _BULLET.match(line)
        if m:
            flush()
            level = min(len(m.group(1).replace("\t", "  ")) // 2, 3)
            ordered = m.group(2)[0].isdigit()
            counters[level] = counters.get(level, 0) + 1 if ordered else 0
            for deeper in [k for k in counters if k > level]:
                counters.pop(deeper)
            blocks.append(("bullet", level, ordered, counters.get(level, 0), m.group(3).strip()))
            i += 1
            continue
        if stripped.startswith(">"):
            flush()
            blocks.append(("quote", stripped.lstrip(">").strip()))
            i += 1
            continue
        para.append(stripped)
        i += 1
    flush()
    return blocks


def _cells(row: str) -> List[str]:
    return [c.strip() for c in row.strip().strip("|").split("|")]


def inline_runs(text: str) -> List[Tuple[str, Dict[str, Any]]]:
    """``[(text, {bold, italic, code, link})]`` for **bold**, *italic*,
    `code`, [text](url); an image reference becomes its alt text."""
    out: List[Tuple[str, Dict[str, Any]]] = []
    pos = 0
    for m in _INLINE.finditer(text or ""):
        if m.start() > pos:
            out.append((text[pos:m.start()], {}))
        tok = m.group(0)
        if tok.startswith("**"):
            out.append((tok[2:-2], {"bold": True}))
        elif tok.startswith("`"):
            out.append((tok[1:-1], {"code": True}))
        elif tok.startswith("!["):
            alt = tok[2:tok.index("]")]
            if alt:
                out.append((alt, {"italic": True}))
        elif tok.startswith("["):
            label, url = tok[1:tok.index("]")], tok[tok.index("(") + 1:-1]
            out.append((label, {"link": url}))
        else:
            out.append((tok[1:-1], {"italic": True}))
        pos = m.end()
    if pos < len(text or ""):
        out.append((text[pos:], {}))
    return out


def _plain(text: str) -> str:
    return "".join(t for t, _ in inline_runs(text))


# ── text size estimate ───────────────────────────────────────────────────────

def _lines(text: str, size: float, width: float) -> int:
    per_line = max(1, int(width / (size * 0.53)))
    return sum(max(1, math.ceil(len(part) / per_line)) for part in (text or " ").split("\n"))


def blocks_height(blocks: List[Tuple], size: float, width: float) -> float:
    """Estimated rendered height in px of ``blocks`` at font ``size`` px."""
    h = 0.0
    for b in blocks:
        kind = b[0]
        if kind == "heading":
            s = size * (1.3 if b[1] <= 2 else 1.12)
            h += _lines(_plain(b[2]), s, width) * s * 1.2 + s * 0.45
        elif kind == "bullet":
            indent = 30 * (b[1] + 1)
            h += _lines(_plain(b[4]), size, width - indent) * size * 1.34 + size * 0.34
        elif kind in ("para", "quote"):
            h += _lines(_plain(b[1]), size, width - (24 if kind == "quote" else 0)) * size * 1.3 + size * 0.5
        elif kind == "code":
            s = size * 0.8
            h += (b[1].count("\n") + 1) * s * 1.35 + s
        elif kind == "table":
            s = size * 0.8
            cols = max(len(r) for r in b[1]) or 1
            tallest = [max(_lines(_plain(c), s, width / cols - 16) for c in r) for r in b[1]]
            h += sum(n * s * 1.25 + s * 0.9 for n in tallest) + size * 0.5
        elif kind == "hr":
            h += size * 0.9
    return h


def fit_size(blocks: List[Tuple], width: float, height: float, base: float, minimum: float = 12) -> float:
    size = base
    while size > minimum and blocks_height(blocks, size, width) > height:
        size -= 1
    return size


# ── drawing primitives ───────────────────────────────────────────────────────

def _rgb(hex_color: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string((hex_color or "#000000").lstrip("#")[:6].upper())


def _emu(v: float) -> int:
    return int(round(v * PX))


def _alpha(shape, pct: int) -> None:
    """Give a solid-filled shape ``pct`` percent opacity."""
    from lxml import etree
    from pptx.oxml.ns import qn
    clr = shape._element.spPr.find(qn("a:solidFill"))
    if clr is not None and len(clr):
        a = etree.SubElement(clr[0], qn("a:alpha"))
        a.set("val", str(int(pct * 1000)))


def _rect(slide, x, y, w, h, color, *, radius: float = 0, alpha: int = 100, shape_kind=None):
    from pptx.enum.shapes import MSO_SHAPE
    kind = shape_kind or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    shp = slide.shapes.add_shape(kind, _emu(x), _emu(y), _emu(w), _emu(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(color)
    shp.line.fill.background()
    shp.shadow.inherit = False
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = min(0.5, radius / max(1, min(w, h)))
    if alpha < 100:
        _alpha(shp, alpha)
    return shp


def _oval(slide, x, y, d, color, alpha=100):
    from pptx.enum.shapes import MSO_SHAPE
    return _rect(slide, x, y, d, d, color, alpha=alpha, shape_kind=MSO_SHAPE.OVAL)


def _gradient_bg(slide, c1, c2, angle=35):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, _emu(W), _emu(H))
    shp.fill.gradient()
    shp.fill.gradient_angle = angle
    stops = shp.fill.gradient_stops
    stops[0].color.rgb = _rgb(c1)
    stops[1].color.rgb = _rgb(c2)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _fade(slide, x, y, w, h, color="#000000", bottom_alpha=78, layers=30):
    """Darken towards the bottom so text over a photo stays readable without a
    hard edge.

    Transparency inside a gradient fill is dropped by Keynote and Quick Look,
    which paint it opaque, so the fade is built from solid layers instead: each
    starts a little lower and all run to the bottom, so their alphas compound
    downwards. Nested rather than side by side, there is no seam between two
    anti-aliased edges to show as a hairline.
    """
    from pptx.enum.shapes import MSO_SHAPE
    per_layer = 1 - (1 - bottom_alpha / 100) ** (1 / layers)
    bottom = _emu(y + h)
    for i in range(layers):
        top = _emu(y + h * i / layers)
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _emu(x), top, _emu(w), bottom - top)
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(color)
        shp.line.fill.background()
        shp.shadow.inherit = False
        _alpha(shp, max(1, round(per_layer * 100)))


def _box(slide, x, y, w, h, *, anchor="top"):
    from pptx.enum.text import MSO_ANCHOR
    tb = slide.shapes.add_textbox(_emu(x), _emu(y), _emu(w), _emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                          "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    return tb, tf


def _para(tf, first: bool):
    return tf.paragraphs[0] if first else tf.add_paragraph()


def _runs(p, text: str, size: float, color: str, *, bold=False, italic=False, accent="#1e3a8a"):
    from pptx.util import Pt
    for chunk, fmt in inline_runs(text) or [("", {})]:
        r = p.add_run()
        r.text = chunk
        f = r.font
        f.size = Pt(size * 0.75)
        f.bold = bold or bool(fmt.get("bold"))
        f.italic = italic or bool(fmt.get("italic"))
        f.color.rgb = _rgb(accent if fmt.get("link") else color)
        f.name = "Courier New" if fmt.get("code") else FONT
        if fmt.get("link") and re.match(r"^(https?:|mailto:)", fmt["link"]):
            r.hyperlink.address = fmt["link"]


def _spacing(p, size: float, after: float, line: float = 1.15):
    from pptx.util import Pt
    p.space_after = Pt(after * 0.75)
    p.line_spacing = line


def _bullet(p, level: int, ordered: bool, number: int, size: float, color: str):
    from lxml import etree
    from pptx.oxml.ns import qn
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(_emu(30 * (level + 1))))
    pPr.set("indent", str(-_emu(30)))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color.lstrip("#").upper())
    if ordered:
        au = etree.SubElement(pPr, qn("a:buAutoNum"))
        au.set("type", "arabicPeriod")
        if number > 1:
            au.set("startAt", str(number))
    else:
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•" if level % 2 == 0 else "–")


def _autofit(tf):
    from lxml import etree
    from pptx.oxml.ns import qn
    body = tf._txBody.find(qn("a:bodyPr"))
    for child in list(body):
        if child.tag in (qn("a:spAutoFit"), qn("a:noAutofit"), qn("a:normAutofit")):
            body.remove(child)
    etree.SubElement(body, qn("a:normAutofit"))


def _text(slide, x, y, w, h, text, size, color, *, bold=False, italic=False,
          align="left", anchor="top", accent="#1e3a8a"):
    from pptx.enum.text import PP_ALIGN
    tb, tf = _box(slide, x, y, w, h, anchor=anchor)
    for n, line in enumerate((text or "").split("\n")):
        p = _para(tf, n == 0)
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        _runs(p, line, size, color, bold=bold, italic=italic, accent=accent)
        _spacing(p, size, 0, 1.05)
    return tb


def _markdown(slide, x, y, w, h, md, t, accent, *, base=26, minimum=13, color=None, muted=None):
    """Lay ``md`` out in the box, text and tables in document order, at the
    largest size that fits. Returns the size used."""
    blocks = parse_markdown(md)
    if not blocks:
        return base
    size = fit_size(blocks, w, h, base, minimum)
    color = color or t["text"]
    muted = muted or t["muted"]
    cursor = y
    segment: List[Tuple] = []

    def flush_text():
        nonlocal cursor, segment
        if not segment:
            return
        seg_h = blocks_height(segment, size, w)
        _, tf = _box(slide, x, cursor, w, max(seg_h, size * 1.4))
        _autofit(tf)
        for n, b in enumerate(segment):
            p = _para(tf, n == 0)
            kind = b[0]
            if kind == "heading":
                s = size * (1.3 if b[1] <= 2 else 1.12)
                _runs(p, b[2], s, color, bold=True, accent=accent)
                _spacing(p, s, s * 0.45, 1.05)
            elif kind == "bullet":
                _bullet(p, b[1], b[2], b[3], size, accent)
                _runs(p, b[4], size, color, accent=accent)
                _spacing(p, size, size * 0.32)
            elif kind == "quote":
                _runs(p, b[1], size, muted, italic=True, accent=accent)
                _spacing(p, size, size * 0.5)
            elif kind == "code":
                from pptx.util import Pt
                for m, line in enumerate(b[1].split("\n")):
                    if m:
                        p = tf.add_paragraph()
                    r = p.add_run()
                    r.text = line or " "
                    r.font.name = "Courier New"
                    r.font.size = Pt(size * 0.8 * 0.75)
                    r.font.color.rgb = _rgb(color)
                    _spacing(p, size, 0, 1.1)
                p.space_after = Pt(size * 0.6 * 0.75)
            elif kind == "hr":
                _runs(p, "", size * 0.5, muted)
                _spacing(p, size, size * 0.4)
            else:
                _runs(p, b[1], size, color, accent=accent)
                _spacing(p, size, size * 0.5)
        cursor += seg_h
        segment = []

    for b in blocks:
        if b[0] == "table":
            flush_text()
            cursor += size * 0.4
            cursor += _table(slide, x, cursor, w, b[1], size * 0.8, t, accent) + size * 0.5
        else:
            segment.append(b)
    flush_text()
    return size


def _table(slide, x, y, w, rows, size, t, accent) -> float:
    from pptx.util import Pt
    cols = max(len(r) for r in rows)
    tallest = [max(_lines(_plain(c), size, w / cols - 16) for c in r) for r in rows]
    heights = [n * size * 1.25 + size * 0.9 for n in tallest]
    total = sum(heights)
    gt = slide.shapes.add_table(len(rows), cols, _emu(x), _emu(y), _emu(w), _emu(total))
    tbl = gt.table
    tbl.first_row = True
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = _emu(heights[ri])
        for ci in range(cols):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(accent if ri == 0 else (t["surface"] if ri % 2 else t["bg"]))
            cell.margin_left = cell.margin_right = _emu(8)
            cell.margin_top = cell.margin_bottom = _emu(4)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            text = row[ci] if ci < len(row) else ""
            _runs(p, text, size, t["heroText"] if ri == 0 else t["text"], bold=ri == 0, accent=accent)
            for r in p.runs:
                r.font.size = Pt(size * 0.75)
    return total


# ── images ───────────────────────────────────────────────────────────────────

_PPTX_IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}


def _image_source(path: Path):
    """``(file or stream, (width, height))`` python-pptx can embed; a format it
    cannot (webp, avif) is converted to PNG in memory."""
    from PIL import Image
    with Image.open(path) as img:
        size = img.size
        if path.suffix.lower() in _PPTX_IMAGE_TYPES:
            return str(path), size
        buf = io.BytesIO()
        img.convert("RGBA").save(buf, format="PNG")
        buf.seek(0)
        return buf, size


def _cover(slide, src, size, x, y, w, h):
    pic = slide.shapes.add_picture(src, _emu(x), _emu(y), _emu(w), _emu(h))
    iw, ih = size
    box, img = w / h, iw / max(1, ih)
    if img > box:
        c = (1 - box / img) / 2
        pic.crop_left = pic.crop_right = c
    elif img < box:
        c = (1 - img / box) / 2
        pic.crop_top = pic.crop_bottom = c
    return pic


# ── the deck ─────────────────────────────────────────────────────────────────

def ordered_slides(spec: Dict[str, Any]) -> List[Dict[str, Any]]:
    coll = spec.get("slides") or {}
    items = ([{"id": k, **(v or {})} for k, v in coll.items()] if isinstance(coll, dict)
             else [{"id": str(i), **(v or {})} for i, v in enumerate(coll)])
    return sorted(items, key=lambda s: s.get("order", 0))


def theme_for(spec: Dict[str, Any]) -> Dict[str, str]:
    t = dict(SLIDE_THEMES.get(spec.get("theme") or "light") or SLIDE_THEMES["light"])
    if spec.get("accent"):
        t["accent"] = spec["accent"]
    return t


def pptx_filename(title: str) -> str:
    slug = re.sub(r"[^\w\-]+", "-", title or "", flags=re.UNICODE).strip("-_")[:60]
    return f"{slug or 'slides'}.pptx"


def build_pptx(view: Dict[str, Any], *, with_warnings: bool = False,
               resolve_asset: Optional[Callable[[str], Optional[Path]]] = None):
    """The .pptx bytes for a slides ``view`` (the stored envelope).

    With ``with_warnings`` returns ``(bytes, warnings)``, the warnings naming
    anything the file could not carry (a remote image, an unreadable asset).
    """
    from pptx import Presentation

    spec = view.get("spec") or {}
    t = theme_for(spec)
    view_id = view.get("view_id") or ""
    if resolve_asset is None:
        from views.store import view_asset_path

        def resolve_asset(name: str) -> Optional[Path]:
            return view_asset_path(view_id, name) if view_id else None

    prs = Presentation()
    prs.slide_width, prs.slide_height = _emu(W), _emu(H)
    prs.core_properties.title = view.get("title") or "Slides"
    blank = prs.slide_layouts[6]
    warnings: List[str] = []
    slides = ordered_slides(spec)
    section_no = 0

    for number, s in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank)
        layout = s.get("layout") or "content"
        accent = s.get("accent") or t["accent"]
        if layout == "section":
            section_no += 1
        ctx = _Ctx(slide, s, t, accent, number, len(slides), spec, resolve_asset, warnings, section_no)
        _LAYOUTS.get(layout, _content)(ctx)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]

    buf = io.BytesIO()
    prs.save(buf)
    data = buf.getvalue()
    return (data, warnings) if with_warnings else data


class _Ctx:
    def __init__(self, slide, s, t, accent, number, total, spec, resolve, warnings, section_no):
        self.slide, self.s, self.t, self.accent = slide, s, t, accent
        self.number, self.total, self.spec = number, total, spec
        self.resolve, self.warnings, self.section_no = resolve, warnings, section_no

    def bg(self, color=None):
        fill = self.slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(color or self.t["bg"])

    def image(self, x, y, w, h) -> bool:
        ref = (self.s.get("image") or "").strip()
        name = ref[len("asset://"):] if ref.startswith("asset://") else ref
        if re.match(r"^(https?:|data:)", name, re.I):
            self.warnings.append(f"slide {self.number}: remote image not embedded ({name[:80]})")
            self._placeholder(x, y, w, h, name)
            return False
        path = self.resolve(name) if name else None
        if path is None:
            if name:
                self.warnings.append(f"slide {self.number}: image asset not found ({name})")
            self._placeholder(x, y, w, h, name)
            return False
        try:
            src, size = _image_source(path)
            _cover(self.slide, src, size, x, y, w, h)
            return True
        except Exception as exc:  # noqa: BLE001 - one bad image never fails the deck
            self.warnings.append(f"slide {self.number}: image {name} could not be embedded ({exc})")
            self._placeholder(x, y, w, h, name)
            return False

    def _placeholder(self, x, y, w, h, label):
        _rect(self.slide, x, y, w, h, self.t["surface"])
        _text(self.slide, x + 24, y + h / 2 - 20, w - 48, 40, f"🖼 {label}" if label else "🖼",
              16, self.t["muted"], align="center", anchor="middle")

    def header(self, x=80, w=1120, y=56, dark=False) -> float:
        """Icon, title, accent bar, subtitle; returns where the body starts."""
        s, t = self.s, self.t
        title_color = t["heroText"] if dark else t["text"]
        tx = x
        if s.get("icon"):
            _text(self.slide, x, y, 60, 60, s["icon"], 40, title_color, anchor="middle")
            tx += 64
        title = s.get("title") or ""
        tsize = 40 if len(title) <= 48 else 32
        _text(self.slide, tx, y, w - (tx - x), 64, title, tsize, title_color, bold=True, anchor="middle")
        _rect(self.slide, x, y + 76, 64, 6, self.accent, radius=3)
        top = y + 100
        if s.get("subtitle"):
            _text(self.slide, x, top, w, 34, s["subtitle"], 22, t["heroMuted"] if dark else t["muted"])
            top += 44
        return top

    def footer(self, x=80, w=1120, dark=False):
        color = self.t["heroMuted"] if dark else self.t["muted"]
        if self.spec.get("footer"):
            _text(self.slide, x, 676, w - 120, 24, self.spec["footer"], 13, color)
        if self.spec.get("numbers", True):
            _text(self.slide, x + w - 110, 676, 110, 24, f"{self.number} / {self.total}", 13, color, align="right")


def _title(c: _Ctx):
    s, t = c.s, c.t
    _gradient_bg(c.slide, t["hero"], t["hero2"])
    _oval(c.slide, 880, -180, 600, t["accent2"], alpha=22)
    _oval(c.slide, 1040, 450, 380, c.accent, alpha=35)
    y = 170
    if s.get("icon"):
        _text(c.slide, 96, y, 100, 80, s["icon"], 60, t["heroText"])
    _rect(c.slide, 96, 262, 88, 8, t["accent2"], radius=4)
    title = s.get("title") or ""
    size = 60 if len(title) <= 40 else 46
    title_h = min(3, _lines(title, size * 1.05, 900)) * size * 1.15
    _text(c.slide, 96, 284, 900, title_h, title, size, t["heroText"], bold=True)
    y = 284 + title_h + 22
    if s.get("subtitle"):
        _text(c.slide, 96, y, 900, 70, s["subtitle"], 26, t["heroMuted"])
        y += 26 * 1.3 * min(2, _lines(s["subtitle"], 26, 900)) + 26
    if s.get("body"):
        _markdown(c.slide, 96, y, 900, max(60, 650 - y), s["body"], t, t["accent2"], base=20, minimum=12,
                  color=t["heroMuted"], muted=t["heroMuted"])
    if c.spec.get("footer"):
        _text(c.slide, 96, 672, 900, 24, c.spec["footer"], 14, t["heroMuted"])


def _section(c: _Ctx):
    s, t = c.s, c.t
    c.bg(c.accent)
    _oval(c.slide, 900, 260, 620, t["heroText"], alpha=8)
    _text(c.slide, 96, 180, 400, 130, f"{c.section_no:02d}", 110, t["heroMuted"], bold=True)
    title = s.get("title") or ""
    title_h = min(3, _lines(title, 54 * 1.05, 1040)) * 54 * 1.15
    _text(c.slide, 96, 320, 1040, title_h, title, 54, t["heroText"], bold=True)
    if s.get("subtitle"):
        _text(c.slide, 96, 320 + title_h + 18, 1040, 80, s["subtitle"], 24, t["heroMuted"])
    if s.get("icon"):
        _text(c.slide, 1080, 120, 120, 110, s["icon"], 80, t["heroText"], align="right")


def _content(c: _Ctx):
    c.bg()
    top = c.header()
    _markdown(c.slide, 80, top + 8, 1120, 650 - top, c.s.get("body") or "", c.t, c.accent)
    c.footer()


def _two_column(c: _Ctx):
    c.bg()
    top = c.header()
    cols = (c.s.get("columns") or ["", ""])[:2]
    h = 650 - top
    blocks = [parse_markdown(col) for col in cols]
    size = min(fit_size(b, 530, h, 24, 12) for b in blocks)
    _rect(c.slide, 639, top + 8, 2, h - 10, c.t["border"])
    for i, col in enumerate(cols):
        _markdown(c.slide, 80 + i * 600, top + 8, 520, h, col, c.t, c.accent, base=size, minimum=size)
    c.footer()


def _image_side(c: _Ctx, image_left: bool):
    c.bg()
    ix = 0 if image_left else 720
    c.image(ix, 0, 560, 720)
    if c.s.get("caption"):
        _rect(c.slide, ix, 660, 560, 60, "#000000", alpha=55)
        _text(c.slide, ix + 20, 668, 520, 44, c.s["caption"], 15, "#ffffff", anchor="middle")
    tx = 620 if image_left else 80
    top = c.header(x=tx, w=580, y=72)
    _markdown(c.slide, tx, top + 8, 580, 650 - top, c.s.get("body") or "", c.t, c.accent, base=24)
    c.footer(x=tx, w=580)


def _image_left(c: _Ctx):
    _image_side(c, True)


def _image_right(c: _Ctx):
    _image_side(c, False)


def _image_full(c: _Ctx):
    s = c.s
    c.bg("#000000")
    c.image(0, 0, W, H)
    _fade(c.slide, 0, 280, W, 440)
    if s.get("icon"):
        _text(c.slide, 80, 400, 80, 60, s["icon"], 44, "#ffffff")
    _text(c.slide, 80, 460, 1120, 110, s.get("title") or "", 48, "#ffffff", bold=True, anchor="bottom")
    if s.get("subtitle"):
        _text(c.slide, 80, 580, 1120, 50, s["subtitle"], 24, "#e5e7eb")
    if s.get("caption"):
        _text(c.slide, 680, 676, 520, 24, s["caption"], 13, "#d1d5db", align="right")


def _quote(c: _Ctx):
    s, t = c.s, c.t
    c.bg()
    _rect(c.slide, 0, 0, 16, H, c.accent)
    if s.get("title"):
        _text(c.slide, 170, 84, 960, 30, s["title"].upper(), 16, c.accent, bold=True)
    _text(c.slide, 64, 150, 110, 150, "“", 150, c.accent, bold=True)
    text = s.get("body") or ""
    size = 40
    while size > 20 and _lines(_plain(text), size, 960) * size * 1.3 > 330:
        size -= 2
    _text(c.slide, 170, 160, 960, 340, _plain(text), size, t["text"], italic=True, anchor="middle")
    if s.get("subtitle"):
        _text(c.slide, 170, 520, 960, 40, f"— {s['subtitle']}", 22, t["muted"])
    c.footer()


def _stats(c: _Ctx):
    s, t = c.s, c.t
    c.bg()
    top = c.header()
    if s.get("body"):
        _markdown(c.slide, 80, top + 4, 1120, 70, s["body"], t, c.accent, base=20, minimum=14)
        top += 80
    items = s.get("items") or []
    n = len(items)
    gap = 32
    w = (1120 - gap * (n - 1)) / n
    h = min(330, 640 - top - 10)
    for i, it in enumerate(items):
        x = 80 + i * (w + gap)
        _rect(c.slide, x, top + 16, w, h, t["surface"], radius=18)
        _rect(c.slide, x, top + 16, 8, h, c.accent)
        value = it.get("value") or ""
        vsize = 72 if len(value) <= 5 else (56 if len(value) <= 8 else 40)
        _text(c.slide, x + 32, top + 44, w - 56, 96, value, vsize, c.accent, bold=True, anchor="middle")
        label = ((it.get("icon") + " ") if it.get("icon") else "") + (it.get("title") or "")
        _text(c.slide, x + 32, top + 150, w - 56, 60, label, 22, t["text"], bold=True)
        if it.get("text"):
            _markdown(c.slide, x + 32, top + 214, w - 56, h - 214, it["text"], t, c.accent,
                      base=17, minimum=11, color=t["muted"])
    c.footer()


def _cards(c: _Ctx):
    s, t = c.s, c.t
    c.bg()
    top = c.header()
    if s.get("body"):
        _markdown(c.slide, 80, top + 4, 1120, 60, s["body"], t, c.accent, base=20, minimum=14)
        top += 66
    items = s.get("items") or []
    n = len(items)
    cols = 2 if n == 4 else min(n, 3)
    rows = math.ceil(n / cols)
    gap = 24
    w = (1120 - gap * (cols - 1)) / cols
    h = (650 - top - 16 - gap * (rows - 1)) / rows
    for i, it in enumerate(items):
        x = 80 + (i % cols) * (w + gap)
        y = top + 16 + (i // cols) * (h + gap)
        _rect(c.slide, x, y, w, h, t["surface"], radius=16)
        _rect(c.slide, x, y, w, 6, c.accent)
        cy = y + 24
        if it.get("icon"):
            _text(c.slide, x + 24, cy, 60, 48, it["icon"], 34, t["text"])
            cy += 54
        if it.get("title"):
            _text(c.slide, x + 24, cy, w - 48, 34, it["title"], 22 if w > 300 else 19, t["text"], bold=True)
            cy += 40
        if it.get("value"):
            _text(c.slide, x + 24, cy, w - 48, 30, it["value"], 18, c.accent, bold=True)
            cy += 32
        if it.get("text"):
            _markdown(c.slide, x + 24, cy, w - 48, y + h - cy - 16, it["text"], t, c.accent,
                      base=17, minimum=11, color=t["muted"])
    c.footer()


def _timeline(c: _Ctx):
    s, t = c.s, c.t
    c.bg()
    top = c.header()
    if s.get("body"):
        _markdown(c.slide, 80, top + 4, 1120, 60, s["body"], t, c.accent, base=20, minimum=14)
        top += 66
    items = s.get("items") or []
    n = len(items)
    seg = 1120 / n
    line_y = top + max(110, (640 - top) / 2 - 50)
    _rect(c.slide, 80, line_y - 2, 1120, 4, c.accent, alpha=35)
    for i, it in enumerate(items):
        cx = 80 + seg * (i + 0.5)
        if it.get("value"):
            _text(c.slide, cx - seg / 2 + 6, line_y - 70, seg - 12, 40, it["value"], 18 if n <= 5 else 15,
                  c.accent, bold=True, align="center", anchor="bottom")
        _oval(c.slide, cx - 12, line_y - 12, 24, c.accent)
        _oval(c.slide, cx - 5, line_y - 5, 10, t["bg"])
        label = ((it.get("icon") + " ") if it.get("icon") else "") + (it.get("title") or "")
        _text(c.slide, cx - seg / 2 + 8, line_y + 28, seg - 16, 60, label, 20 if n <= 4 else 16,
              t["text"], bold=True, align="center")
        if it.get("text"):
            _text(c.slide, cx - seg / 2 + 8, line_y + 72, seg - 16, 640 - line_y - 72,
                  _plain(it["text"]), 16 if n <= 4 else 13, t["muted"], align="center")
    c.footer()


_LAYOUTS: Dict[str, Callable[[_Ctx], None]] = {
    "title": _title, "section": _section, "content": _content, "two_column": _two_column,
    "image_left": _image_left, "image_right": _image_right, "image_full": _image_full,
    "quote": _quote, "stats": _stats, "cards": _cards, "timeline": _timeline,
}

__all__ = ["build_pptx", "pptx_filename", "parse_markdown", "inline_runs", "ordered_slides", "theme_for"]
